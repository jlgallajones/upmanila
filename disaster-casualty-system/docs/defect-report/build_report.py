"""Build the SCRIBER defect and resolution report from the retained issue history."""
from pathlib import Path
from collections import Counter
import importlib.util
import json
import re
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/docx/SCRIBER-Defect-and-Resolution-Report.docx'
QA=ROOT/'tmp/defect-report'

# Each record preserves source references while resolving obsolete descriptions.
FINDINGS=[
('Authentication and session recovery','7 10A 10B 33','Defect','Recorded fix',
'Mobile record views produced unauthorized responses and could fail when history data omitted startTime. Password recovery was a placeholder, and an accidental logout click immediately ended the session. A separate investigation also treated HTTP 304 as a fetch error without checking whether cached data was usable.',
'The recorded changes strengthen session handling and defensive access to optional history fields, add an API-backed recovery email and reset-password screen, and require confirmation before clearing the web session. Cache-validation responses are assessed together with the returned application state; a 304 response alone is not evidence that retrieval failed.',
'Exercise token expiry, recovery links, missing history values and logout cancellation. Confirm that genuinely unauthorized requests remain rejected and that usable cached records still render.'),
('Account roles and deployment compatibility','20 40 41 78 79 102','Defect and compatibility change','Source reviewed',
'The original account model used a generic responder role. Later separation of field, stabilization and facility workflows exposed stale enum values, old API validation and cached mobile profiles. Account creation and exports could fail when the database, backend and client used different role definitions.',
'The current export controller includes field_responder, sa_responder, minor_incident_documenter, documenter and supported legacy roles. Earlier guidance restricting exports to responder and documenter is obsolete. Role migrations and API deployment must accompany the client rollout. Mobile profile refresh removes stale function selection for fixed-role accounts, while the dedicated Minor Incident Documenter flow preserves its longer reporting path.',
'Create, edit, export and sign in with each supported role after migrations are applied. Refresh an account changed from a legacy role and verify its workflow without relying on cleared device storage.'),
('Server enforcement of record visibility','12 27','Security assurance','Verification required',
'Operational users were able to see records outside their intended account or role scope in mobile lists. The history also identifies a broader risk: hiding records in the interface does not establish that direct API requests enforce the same boundary.',
'Account filtering was added to the mobile record view, and the backend contains authentication, role and unit-scope controls. The issue should remain an authorization verification item until list, detail, attachment, history, analytics and export endpoints have been tested using accounts from different units and operational roles.',
'Use a cross-account request matrix against the API. Requests outside the caller’s scope must return no unauthorized data, including nested histories, linked cases and attachment access.'),
('Shared device isolation of offline records','111','Security defect','Source reviewed',
'The offline queue was stored at device level without an owner on each item. Different accounts sharing a device could therefore see another account’s unsynchronized casualty records. An unrestricted queue panel could also prevent access to the submitted-record list.',
'Queue items now carry ownerUserId. Reads, counts, incident assignment, code correction and synchronization select the current owner, while unidentified legacy queue items remain stored without being exposed to another account. The queue panel has its own bounded scroll area and can be collapsed. The current queue implementation contains these ownership filters.',
'Queue work under two accounts on one device. Verify isolation for every queue action, including after logout, session expiry and a legacy queue upgrade. Confirm that a long queue does not block the submitted list.'),
('Offline entry without a cached incident','19 21 35','Data capture defect','Recorded fix',
'Incident selection originally depended on an online API response. A cached list addressed repeat offline use, but a first launch in airplane mode still blocked casualty entry. Simply removing the incident requirement would leave a queued record that the server could not accept.',
'The revised flow uses cached official incidents when available and permits a local placeholder when the device has no saved incident list. A placeholder record retains an optional temporary incident name and is marked as needing assignment. An official active incident must be assigned before synchronization. This supersedes the earlier statement that first-time offline entry always requires a previously cached incident.',
'Test first-launch airplane mode, cached offline entry and reconnection. An unassigned record must remain recoverable locally and must not be submitted with a fabricated incident identifier.'),
('Duplicate identifiers and submission retries','24 36 94 110','Data integrity defect','Source reviewed',
'Offline identifiers could collide with existing records, and generic duplicate messages concealed the actual conflict. A lost response or rapid repeated tap could also leave the user unsure whether a report had been created. Earlier notes describe several different victim-code scopes.',
'The current ensureUniqueIdNumber check evaluates id_number within the selected incident. Responder sequencing is also incident-aware. The submission changes retain a client record identifier across retries, guard overlapping taps and treat an already-created client record as an idempotent success. Generated offline identifiers have a controlled recovery path; manually entered identifiers are not silently replaced. The older account-and-role scope description must not be treated as the current API rule.',
'Run simultaneous submissions and lost-response retries. Confirm one server record per client identifier, a clear conflict message for a duplicate incident code, and explicit user confirmation where visible codes are corrected.'),
('Submission latency and queue throughput','112','Performance defect','Source reviewed',
'A successful core record save was not acknowledged until responder-code refresh and photo uploads completed. Queue synchronization processed reports sequentially, so larger queues appeared unresponsive even while work was progressing.',
'The recorded change acknowledges the core save when the API returns and handles photo uploads and code refresh afterward. Attachment failures are reported as follow-up warnings. The current queue sets a concurrency limit of three independent submissions. This improves throughput without changing the requirement to isolate accounts and preserve per-record retry state.',
'Measure core-save acknowledgement separately from attachment completion. Synchronize at least six queued reports and verify progress, correct ownership, safe retries and no duplicate server records. No benchmark improvement is claimed without timings.'),
('Responder safety consistency during synchronization','56 88 114','Data integrity defect','Recorded fix',
'Offline safety fields could be entered repeatedly, and later synchronization could overwrite the original safety response. Call-down and analytics views also relied on different safety sources, producing inconsistent summaries.',
'The queue now retains safety status, arrival time, protective-equipment time and responder context, and submits the safety response before casualty creation. The recorded backend rule preserves the first valid incident-and-responder response. Call-down records incorporate responder safety data, and Incident Management presents a summary rather than a second independent safety form.',
'Repeat safety submission online and offline and verify that the first entry remains unchanged. Compare the call-down roster, responder summary and analytics after synchronization, including an account with a safety response but no pre-existing call-down row.'),
('Assisted triage calculation and system availability','8B 82 90 91 113 115','Calculation defect','Source reviewed',
'Some supported systems were missing from selectors, some tertiary assessments returned unknown, and calculation order produced incorrect stored categories for particular answer combinations. Mobile and API implementations could disagree. Systems without a configured calculator could block entry instead of accepting a manual final category.',
'The recorded work expands supported primary and secondary selections, adds or corrects assisted logic, aligns mobile and API behavior, and provides manual final-category entry for systems without an automatic calculator. Specific corrections cover METTS, tertiary decision branches, SORT component scoring and the STIEVE special-population override. Current API tests include STIEVE and SIEVE regression cases.',
'Run the software regression suite and compare identical inputs across mobile and API. Obtain operational approval of the reference cases before release. Passing calculation tests does not establish that every configured protocol has been clinically validated.'),
('Incident control of triage mode','103','Workflow and validation change','Recorded fix',
'Responders could independently switch between assisted and unassisted entry, allowing inconsistent assessment modes within one official incident.',
'Triage mode is now an incident setting. The mobile form displays the configured mode, clears incompatible answers when the incident changes, and the API validates the mode during casualty and assessment writes. Manual-only systems retain manual final-category entry. The add-incident-triage-mode.sql migration is a deployment dependency; legacy incidents use the default assisted setting.',
'Test incident changes with an in-progress assessment and reject payloads that conflict with the incident setting. Verify legacy incidents and manual-only systems after the migration.'),
('Role specific entry and navigation','4 25 26 69 71 72 74 81 98','Workflow defect','Recorded fix',
'Entry forms exposed personal details before identification was selected, repeated hospital information, used ambiguous labels and treated Back and Previous as equivalent actions. Incident context and responder identifiers could be difficult to follow through a long form.',
'The revised forms conditionally show identification fields, retain role-specific capture requirements, separate workflow exit from previous-step navigation, and keep incident or casualty context visible. Field triage stage uses the requested two-option control. Facility capture consolidates admission and disposition fields and preserves the admitted unit. Updated record labels and review navigation make the saved state easier to interpret. Earlier restrictions on available systems have been superseded by later triage changes.',
'Walk through identified and unidentified cases for each role. Change incident, stage and admission choice, navigate backward, and reopen the record. Hidden or incompatible values must not leak into the saved payload.'),
('Record details and controlled editing','2 6 23 30 31 32 83 85 89','Record workflow defect','Recorded fix',
'Record details omitted saved assessment answers, lacked latest triage and transport summaries, or showed unrelated role sections. Treatment history returned 404. Rejected-record edits did not clearly return the report to review, and facility Management and Disposition were not accessible for permitted post-submission edits.',
'The recorded corrections align detail sections with role-specific forms, restore treatment retrieval, serialize assessment answers and enrich report summaries. Rejected corrections return to review with visible feedback. Transport and permitted facility sections receive focused edit actions, while protected sections remain locked. Verification labels are normalized, duplicate badges are removed, and the stored minimal category is displayed as Minor where required.',
'Compare submitted form values with API detail, mobile detail, web review and export. Verify permitted and blocked edits, treatment history retrieval and the full reject-correct-resubmit-approve cycle.'),
('Notifications and record navigation','3 34','Defect and enhancement','Partial',
'Responders did not reliably receive rejection notifications. Successful-submission notifications and click-through to the correct casualty record were also identified as incomplete.',
'Rejection notifications are recorded as connected to the submitted record. A success notification must be created only after a valid server record identifier is returned, addressed to the submitting account and linked to the standard record-detail route. The available notes do not establish that the full success-notification path has passed acceptance, so it remains separate from the rejection correction.',
'Verify recipient, unread count and record destination for rejection and successful creation. Failed, retried and offline-only submissions must not create misleading success notifications.'),
('Attachment capture and visibility','13 42 109','Attachment defect','Recorded fix',
'Camera actions opened file selection, uploaded images were not visible in web review, and edit mode displayed only newly selected files. Photo capture was not available consistently across operational roles, weakening case-matching evidence.',
'Capture and import are separate actions. Existing upload, offline payload and preview behavior is reused across role-specific photo entry points. Web record views expose attachments, and mobile editing loads the existing list with thumbnails, names, type and an empty state. Adding a photo does not imply that older attachments should be replaced.',
'Capture and import on a real phone, submit online and offline, then inspect the same attachments from mobile editing and web review. Verify that a failed upload is recoverable without duplicating the casualty.'),
('Draft persistence and save feedback','51 68','Data loss prevention','Recorded fix',
'Long forms had no dependable resume path, and mobile draft saves did not provide clear feedback in the PWA runtime. Users could mistake a saved draft for a submitted record or lose work while starting another entry.',
'Web forms use form_drafts through the drafts API, while mobile casualty forms retain local drafts. Save, resume, update and delete actions are available, and the active form is cleared after a successful save. Account draft payloads exclude temporary passwords. Mobile save confirmation uses the in-app feedback modal and directs the user to Drafts.',
'Save and resume every supported form, including after restart. Confirm the correct owner, exact field restoration, no stored password, and clear separation between a draft and a server-submitted casualty.'),
('Case matching and combined record review','43 44 45 46 47 48 93','Workflow and data linkage change','Source reviewed',
'Separate operational reports for one casualty had no formal link. Initial matching screens allowed ambiguous selection, and opening a linked record still displayed only its own data. The original requirement for all three role records later blocked valid field-and-advanced-medical matches.',
'casualty_case_links connects existing reports without merging or deleting them. Role-specific pickers enforce one record per slot and exclude records already assigned elsewhere. Current API validation requires Field Responder and Advanced Medical Responder records; HCFD is optional. Completed links remain locked, and combined review loads each role section from its corresponding linked report. Older three-mandatory-role descriptions are superseded.',
'Accept a valid two-role match and an optional three-role match. Reject duplicate roles and already-linked records. Reopen every linked report and verify the same case association and correctly sourced role sections.'),
('Analytics source mapping and empty datasets','8 8A 53 86 95 96','Reporting defect','Recorded fix',
'Incident metrics and facility charts could be empty despite saved records because the analytics query omitted assessment_answers and relied on one normalized category. Timelines lacked the earliest facility activation event, and blank charts implied broken reporting.',
'Analytics now selects and normalizes category values from saved assessment and result fields, uses appropriate facility assessments, and builds incident-specific metrics. The recorded updates add facility activation timing, triage comparison summaries and operational series. Empty pie and line datasets are omitted. Reported percentages use the documented denominator, and assisted and unassisted comparisons remain distinguishable.',
'Reconcile each displayed total with source records for a selected incident. Cover alternate stored ESI formats, unknown categories, zero records, linked role reports and multiple incidents. Explicitly verify whether each metric counts reports or distinct victims.'),
('Incident situation reports and downloads','1 57 59 60 97','Reporting defect','Recorded fix',
'Situation reports appeared identical, were difficult to read and could represent a responder-role subset instead of the selected incident. Long location values overlapped, and filenames and draft or final actions were inconsistent.',
'The recorded report generator uses incident-wide data, includes the analytics snapshot and role coverage, and lays out wrapped metadata on separate lines. Draft and final actions are distinct; final generation remains dependent on incident closure. Downloads use the office, incident and operation date in the filename. Earlier responder-function filtering is superseded by the incident-wide requirement.',
'Export two incidents with different records and compare scope, totals and filenames. Test long locations, empty chart data, draft output before closure and final output after closure. Confirm the selected incident is visible before generation.'),
('Record filters and historical incident review','11 22 29 39 54 65 75 84','Defect and enhancement','Partial',
'Lists lacked reliable date, role, status and submission-order filters. Long timeline text impaired readability, closed incidents were difficult to reopen for review, and ESI filters assumed one stored value format.',
'The recorded corrections add date and account filters, explicit apply and clear actions, newest or oldest ordering, wrapped timeline text and time-based sorting. Incident History includes closed records and opens incident-scoped records or analytics. ESI matching normalizes supported stored formats. The broader proposed set of clinical, admission and local-sync filters remains a scope item unless separately accepted.',
'Combine and clear filters against a known dataset. Test closed incidents, equal timestamps, long event text and alternate ESI values. Confirm the accepted filter inventory rather than assuming every filter proposed in the early log exists.'),
('Incident lifecycle and management controls','38 52 55 87 92','Workflow and protection change','Recorded fix',
'Management sections were ordered inconsistently, a unit history action opened the wrong destination, and closed incidents still exposed editable controls. Summary cards showed raw field names, while the obsolete evacuation-center workspace remained visible.',
'The recorded updates reorganize management sections and the call-down roster, route unit history to analytics and format summary labels. Closure can proceed with an explicit warning for missing details, and reopening follows an approval workflow. Closed or archived incidents render view-only controls with client-side guards against stale edits. The standalone evacuation-center UI was removed while the underlying data and API were preserved.',
'Verify closure, post-closure export, reopen requests and approved reopening. Attempt a stale form save and a direct API write against a closed incident; interface locking alone is insufficient evidence of server protection.'),
('Temporary incident command delegation','106 107 108','Authorization and workflow change','Source reviewed',
'Shift handovers needed temporary incident control without replacing permanent unit ownership or granting broad administrator access. A delegated responder had no restricted web workspace or analytics entry point.',
'The recorded design retains incidents.created_by as the permanent owner and uses incident-specific backup and delegation records. Delegation requires a reason, records an audit event, notifies the selected user and restricts the Acting IC workspace to assigned incident operations. The primary owner’s management writes are blocked during active delegation. Analytics is limited to the delegated incident. The delegation migration is required before deployment.',
'Test designation, delegation and resumption with separate accounts. Verify denied primary-owner writes during delegation, denied unrelated-incident requests, continued responder entry, notification delivery and preserved permanent ownership.'),
('Bulk import validation and partial success','15 61 62 63 64','Import and integrity defect','Recorded fix',
'Imports could create records but return a generic error or HTTP 500 when audit logging failed afterward. Incomplete unit-account rows reached account creation, and completion feedback did not distinguish created, skipped and failed rows.',
'Preview and server validation now reject missing required account fields and expose row-level outcomes. Completed imports return a success or warning summary with a visible completion dialog. Audit summaries include an entity identifier, and audit failure after creation no longer changes a completed import into an apparent total failure. The audit failure still requires operational follow-up.',
'Import valid, duplicate and invalid rows together, then simulate post-create audit failure. The response must match actual created accounts, prevent duplicate retries and retain enough evidence to repair missing audit activity.'),
('Audit coverage and attribution','16 50','Traceability defect','Recorded fix',
'Audit coverage initially focused on accounts and did not provide a complete view of mobile submissions, verification, matching, attachments, imports and incident changes. Activity could be attributed outside the expected administrator scope.',
'The recorded changes extend logging across major mobile and web writes and resolve unit visibility through the actor and the administrator responsible for the account. Delegation, resets and imports also require traceable events. Logging a successful operation and authorizing a user to read its audit record are separate responsibilities.',
'Perform one action in each audited workflow and inspect actor, entity, action, time and unit visibility. Reconcile suppressed post-create logging failures with the import recovery procedure.'),
('Interface feedback navigation and refresh','9 10 17 58 70 73 80 99 100 101','Usability defect','Recorded fix',
'Dashboard navigation became difficult to scan, actionable and display-only cards looked alike, and progress feedback was inconsistent. Dark-mode autofill and icon alignment defects reduced login readability, while some screens failed to follow device theme changes.',
'The recorded corrections group navigation, refine collapsed and mobile layouts, distinguish actionable cards, and expose refresh progress and completion states. Realtime refresh was extended to analytics and audit views; subscription health and fallback refresh still need operational checks. Light and dark palettes respect system mode and allow explicit overrides. Autofill, field clipping, icon placement and summary title contrast were corrected.',
'Review narrow and wide screens, keyboard focus, autofill, system theme changes and a manual override. Verify that refresh failure, partial success and completion are visually distinct and do not rely on color alone.'),
('Runtime configuration and API URL normalization','66 67','Configuration defect','Recorded fix',
'Client source contained deployment-specific runtime values and a tracked mobile test-user variable. A quoted API base URL later produced an encoded quote in deployed login requests and a route-not-found response.',
'The recorded changes move dashboard configuration to window.DCMS_CONFIG, retain example files and exclude local runtime files from version control. The mobile client trims wrapping quotes, whitespace and trailing slashes before constructing its API client. Publishable client configuration remains public by design; moving it out of app.js does not make it secret or prove that historical credentials were rotated.',
'Build with clean and incorrectly quoted URLs. Inspect the outgoing route and the final bundle for server-only secrets or test-user data. Verify that an actual secret exposure would receive a separate history and rotation review.'),
('Treatment value normalization and display labels','104 105','Data compatibility change','Recorded fix',
'Scooter and Play and Run represented the same requested treatment strategy, while Field Responder displays used an abbreviated Scoop and Run label. Inconsistent values risked fragmenting stored data and analytics.',
'The revised treatment mapping uses scooter, normalizes legacy play_and_run values and supplies a migration to update existing rows and constraints. The user-facing label becomes Scoop and Run/Load and Go while its internal scoop_and_run value remains unchanged. Data normalization and display wording are deliberately handled separately.',
'Test legacy rows, new selections, payloads, record details and analytics after the migration. Confirm that previously stored values remain readable and do not create duplicate strategy buckets.'),
('Exports backups and destructive operation safeguards','14 18','Administration and recovery change','Recorded fix',
'Administrators lacked complete export and profile workflows, while delete and reset actions relied on generic browser confirmations with insufficient context about affected records.',
'The recorded changes add profile editing, password change, incident exports and system backup actions. Destructive operations use in-app confirmation, typed input or password verification where applicable, affected counts and completion summaries. Export scope is enforced through backend permissions, with attachments or attachment references included as appropriate.',
'Verify permission checks for every export and backup endpoint. Compare preview and final reset counts, cancel at each confirmation step, and verify the resulting audit event. A successful download alone does not demonstrate that a backup can be restored.'),
('Submission documentation and operational handover','49 76 77','Documentation delivery','Recorded delivery',
'Technical documentation, user instructions and system diagrams were mixed with development notes, making submission and handover material difficult to review independently.',
'The history records separate technical documentation, user guidance, flowcharts, data-flow diagrams and architecture or schema references. These items are documentation deliverables rather than runtime defects. Current SCRIBER artifacts should use consistent product naming and match the deployed role, matching, offline and reporting rules.',
'Review the documentation against the final client and API behavior, preserve clear setup and migration instructions, and run the demonstration workflow using the documented steps.'),
('Facility selection and reference maintenance','5 28 37','Reference data defect','Recorded fix',
'Facility workflows could stop after incident selection, and free-text receiving-facility values could diverge from the official directory. Facilities could be created but not corrected through a row-level edit action.',
'The recorded changes restore the incident-to-official-facility selection sequence, use stable facility references and provide a scoped facility update endpoint with an edit modal and audit event. Facility identity must remain stable when a display name, contact or active status changes; historical reports should not depend solely on free-text labels.',
'Complete disruption and closure forms through facility selection. Edit and deactivate a facility, then verify account scope, audit history, new selection behavior and readability of existing linked records.'),
]


def make_doc():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    norm=doc.styles['Normal'];norm.font.size=Pt(11)
    norm.paragraph_format.line_spacing=1.08;norm.paragraph_format.space_after=Pt(7)
    title=doc.styles['Title'];title.font.size=Pt(26);title.font.bold=True
    title.paragraph_format.space_after=Pt(12)
    for name,size in [('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
        s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(13);s.paragraph_format.space_after=Pt(6)
    doc.core_properties.title='SCRIBER Defect and Resolution Report'
    doc.core_properties.subject='Defect history corrective actions and release verification'
    doc.core_properties.author='Project Engineering'
    doc.core_properties.last_modified_by='Project Engineering'
    doc.core_properties.comments=''
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run('SCRIBER  |  Engineering report  |  ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    return doc


def paragraph(doc,text,label=None):
    for old,new in [
        ('The recorded changes','The changes'),('The recorded corrections','The corrections'),
        ('The recorded updates','The updates'),('The recorded work','The implementation'),
        ('The recorded change acknowledges','The revised submission flow acknowledges'),
        ('The recorded backend rule','The revised backend rule'),
        ('The recorded report generator','The report generator'),
        ('The recorded design','The delegation design'),
    ]:text=text.replace(old,new)
    p=doc.add_paragraph()
    if label:p.add_run(label+' ').bold=True
    p.add_run(text)
    return p


def matrix(doc,headers,rows,widths):
    table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
    for c,width in zip(table.columns,widths):c.width=Inches(width)
    for cell,value in zip(table.rows[0].cells,headers):cell.text=value
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,value in zip(table.add_row().cells,row):cell.text=str(value)
    for index,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for col,cell in enumerate(row.cells):
            cell.width=Inches(widths[col])
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');mar.append(node)
            pr.append(mar)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8EC' if index==0 else ('F4F6F7' if index%2==0 else 'FFFFFF'));pr.append(shade)
            align=OxmlElement('w:vAlign');align.set(qn('w:val'),'center');pr.append(align)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
                for r in p.runs:r.font.size=Pt(10);r.bold=index==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def main():
    source=(ROOT/'BUGS_AND_SOLUTIONS.md').read_text(encoding='utf-8')
    source_ids=re.findall(r'^## (\d+[A-Z]?)\.',source,re.M)
    covered=[item for finding in FINDINGS for item in finding[1].split()]
    assert len(source_ids)==119 and set(source_ids)==set(covered),set(source_ids)-set(covered)
    assert all(count==1 for count in Counter(covered).values()),'Duplicate trace reference'
    verification=json.loads((ROOT/'output/verification/verification-latest.json').read_text(encoding='utf-8-sig'))
    doc=make_doc();doc.add_paragraph('SCRIBER Defect and Resolution Report','Title')
    doc.add_paragraph('Disaster Casualty Management System','Subtitle')
    paragraph(doc,'22 October 2026  |  Engineering review  |  Version 1')
    doc.add_paragraph('Engineering assessment','Heading 1')
    paragraph(doc,'The defect history records corrective work across mobile submission, account access, incident administration, triage calculations and operational reporting. The latest retained verification run completed 10 checks with no recorded failures on 21 October 2026. That result supports build and selected regression-test stability; it does not close outstanding live acceptance checks.')
    paragraph(doc,'Release acceptance should focus on direct API authorization, shared-device offline isolation, concurrent retries, incident command handover and the consistency of calculations across clients. These areas can affect record confidentiality, integrity or operational decisions even when a build succeeds.')
    paragraph(doc,'This report consolidates 119 numbered history entries into 29 findings. It preserves the original item references, separates corrective work from feature changes, and replaces superseded role, matching and offline descriptions with the latest recorded behavior. Individual findings state the failure, corrective action and evidence needed for closure.')
    paragraph(doc,'FR refers to Field Responder, AMP to Advanced Medical Responder, HCFD to Healthcare Facility Documenter, and Acting IC to Acting Incident Commander. The legacy sa_responder role identifier remains in use for the stabilization or advanced-medical workflow.')
    doc.add_paragraph('Status interpretation','Heading 2')
    matrix(doc,['Status','Meaning'],[
      ('Source reviewed','A targeted source inspection supports the described change. Deployment and end-to-end acceptance still require their own evidence.'),
      ('Recorded fix','The development history records a correction. This report does not assign independent production closure.'),
      ('Partial','A recorded correction exists, but a related acceptance requirement or proposed feature remains unconfirmed.'),
      ('Verification required','The remaining work is assurance of a control or behavior, rather than an assertion that the current implementation is defective.'),
      ('Recorded delivery','A documentation deliverable is recorded separately from application defects.'),
    ],[1.4,5.5])
    doc.add_paragraph('Current rules that supersede earlier notes','Heading 2')
    paragraph(doc,'Account exports support the separated roles and legacy compatibility values. Matching requires Field Responder and Advanced Medical Responder reports, with HCFD optional. First-time offline entry can retain a local placeholder pending official incident assignment. The current ID-number duplicate check is incident-scoped. Situation reports cover the selected incident as a whole.')

    doc.add_page_break();doc.add_paragraph('Defect findings and corrective actions','Heading 1')
    for index,(title,ids,kind,status,problem,correction,acceptance) in enumerate(FINDINGS,1):
        doc.add_paragraph(f'Finding {index:02d} {title}','Heading 2')
        p=paragraph(doc,f'History items {", ".join(ids.split())}  |  {kind}  |  {status}')
        for r in p.runs:r.font.size=Pt(9)
        p.paragraph_format.keep_with_next=True
        paragraph(doc,problem,'Failure and impact')
        paragraph(doc,correction,'Corrective action')
        paragraph(doc,acceptance,'Acceptance check')

    doc.add_page_break();doc.add_paragraph('Verification evidence','Heading 1')
    paragraph(doc,'The retained verification summary identifies a run from 21 October 2026 at 22:47:20 to 22:47:44 UTC+08. It records 10 checks passed and 0 failed. These are existing results, not a new execution performed for this report.')
    matrix(doc,['Recorded check','Result'],[(t['test'],t['result']) for t in verification['tests']],[5.9,1.0])
    paragraph(doc,'The final regression checklist explicitly distinguishes implementation coverage from live testing. Frontend console behavior and unexpected backend errors remain unchecked there. Several workflow items have only partial live confirmation. A checked implementation box must therefore not be used as evidence that every role and device combination has passed.')
    doc.add_paragraph('Targeted source review','Heading 2')
    paragraph(doc,'Inspection confirmed the current export role allowlist, incident-scoped ID-number duplicate lookup, required case-matching roles, account-filtered queue reads and queue concurrency limit. Triage test files contain regression cases for the STIEVE override and SIEVE branches. These checks resolve contradictions in the history; they are not a replacement for deployed endpoint tests.')

    doc.add_paragraph('Release acceptance priorities','Heading 1')
    paragraph(doc,'The following priorities are engineering recommendations based on potential impact. They are not incident severity ratings or commitments to a delivery date.')
    matrix(doc,['Priority','Required evidence','Responsible function'],[
      ('P1','Direct API tests across accounts, roles and units, including attachments, linked cases, exports and delegated incidents.','Backend and QA'),
      ('P1','Shared-device queue tests, expired-session recovery, placeholder incident assignment and repeated or concurrent synchronization.','Mobile and QA'),
      ('P1','Approved reference cases produce consistent saved triage results on mobile and API, including manual-only systems.','Engineering and operational reviewers'),
      ('P1','Acting IC delegation and resumption preserve ownership and enforce write restrictions. Closed incidents reject unauthorized mutations.','Backend and QA'),
      ('P2','Bulk partial-success responses reconcile with created records, and any missing audit activity can be identified and repaired.','Backend and QA'),
      ('P2','Dashboard, case matching, analytics and SitRep totals reconcile against a controlled multi-role incident dataset.','Web and QA'),
      ('P2','Device tests cover camera capture, large photos, failed uploads and storage pressure. Backup contents have a defined restoration test.','Mobile and operations'),
      ('P3','Refactor large form and controller modules, centralize field mapping, and add realtime health or last-updated visibility.','Engineering'),
    ],[.6,4.9,1.4])
    doc.add_paragraph('Regression execution','Heading 2')
    paragraph(doc,'Run the repository verification script from disaster-casualty-system with PowerShell using tools/run-verification.ps1. Preserve its timestamped JSON and log outputs. After automated checks pass, execute the multi-role submission, rejection, resubmission, attachment, offline and incident-handover scenarios against the intended environment.')
    paragraph(doc,'For each acceptance run, retain the application build, migration state, account role and unit, device or browser, input record identifiers, expected result and actual result. Record the reviewer and date when a finding is closed. No per-finding live closure date or severity is inferred from the current notes.')

    doc.add_paragraph('Evidence references','Heading 1')
    refs=[
      ('Development issue history','BUGS_AND_SOLUTIONS.md, numbered items 1 through 115 and supplemental items 8A, 8B, 10A and 10B.'),
      ('Regression status','FINAL_REGRESSION_CHECKLIST.md.'),
      ('Recorded verification run','output/verification/verification-20261021-224720.json and the matching .log file; verification-latest.json points to this run.'),
      ('Verification entry point','tools/run-verification.ps1 and api/package.json.'),
      ('API source reviewed','api/src/controllers/export.controller.ts; casualty.controller.ts; incident.controller.ts; responder-safety-response.controller.ts.'),
      ('Queue and calculation evidence','mobile/src/offline/casualtyQueue.ts; api/src/services/triage/calculate-triage.test.ts.'),
      ('Deployment dependencies','api/sql/add-incident-triage-mode.sql; add-incident-acting-ic-delegation.sql; merge-play-and-run-into-scooter.sql; role migrations under api/sql.'),
    ]
    for name,text in refs:paragraph(doc,text,name)
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    QA.mkdir(parents=True,exist_ok=True)
    # Structural verification complements, but does not substitute for, a render.
    check=Document(OUT)
    text='\n'.join(p.text for p in check.paragraphs)
    assert len([p for p in check.paragraphs if p.text.startswith('Finding ')])==29
    assert 'medsort' not in text.lower()
    assert '5 years' not in text and 'five years' not in text
    with ZipFile(OUT) as z:assert z.testzip() is None
    summary={'output':str(OUT),'source_items':len(source_ids),'findings':len(FINDINGS),'words':len(text.split()),'tables':len(check.tables),'source_coverage':'119 of 119'}
    (QA/'content-check.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps(summary))


if __name__=='__main__':main()
