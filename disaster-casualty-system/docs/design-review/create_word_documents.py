"""Create three separate SCRIBER Word documents from the retained design assets."""
from io import BytesIO
from pathlib import Path
import importlib.util
import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfgen import canvas
import pypdfium2 as pdfium

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
OUT=ROOT/'output/docx'
TMP=ROOT/'tmp/scriber-word'
OUT.mkdir(parents=True,exist_ok=True)
TMP.mkdir(parents=True,exist_ok=True)
spec=importlib.util.spec_from_file_location('design',HERE/'generate_review.py')
design=importlib.util.module_from_spec(spec);spec.loader.exec_module(design)


def figures():
    scenes={'architecture':design.architecture(),'core-erd':design.core_erd(),
            'schema-extensions':design.extensions(),'web-wireframe':design.web_wireframe(),
            'mobile-wireframes':design.mobile_wireframes()}
    for name,scene in scenes.items():
        stream=BytesIO();c=canvas.Canvas(stream,pagesize=(scene.w,scene.h))
        scene.pdf(c,0,0,scene.w,scene.h);c.showPage();c.save()
        pdf=pdfium.PdfDocument(stream.getvalue())
        image=pdf[0].render(scale=2.5).to_pil().convert('RGB')
        image.save(TMP/f'{name}.png')
        if name=='mobile-wireframes':
            for i,label in enumerate(['identity','review','queue']):
                x=30+395*i
                image.crop((int(x*2.5),int(48*2.5),int((x+350)*2.5),int(712*2.5))).save(TMP/f'mobile-{label}.png')


def style_doc(title,subject):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6)
    sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    sec.footer_distance=Inches(.25)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(27);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(9)
    for name,size in [('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
        s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9.5)
    doc.styles['Caption'].paragraph_format.space_after=Pt(8)
    doc.core_properties.title=title;doc.core_properties.subject=subject
    doc.core_properties.author='UP Manila project';doc.core_properties.keywords='SCRIBER'
    doc.add_paragraph(title,'Title')
    p=doc.add_paragraph('Disaster Casualty Management System','Subtitle');p.paragraph_format.space_after=Pt(5)
    p=doc.add_paragraph('22 October 2026  |  For review');p.runs[0].font.size=Pt(9.5)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('SCRIBER  |  ');r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    return doc


def para(doc,text):return doc.add_paragraph(text)
def heading(doc,text):return doc.add_paragraph(text,'Heading 1')
def next_page(doc,title):doc.add_page_break();heading(doc,title)


def picture(doc,name,caption,width=7):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4)
    shape=p.add_run().add_picture(str(TMP/f'{name}.png'),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    doc.add_paragraph(caption,'Caption')


def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=value
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement(f'w:{edge}');node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');margins.append(node)
            props.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement(f'w:{edge}');node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
            props.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8EC' if ri==0 else ('F5F7F8' if ri%2==0 else 'FFFFFF'));props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
                for r in p.runs:r.font.name='Calibri';r.font.size=Pt(10);r.font.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


def save(doc,name):
    path=OUT/name;doc.save(path);print(path)


def architecture_doc():
    doc=style_doc('SCRIBER System Architecture','System components and operational data flow')
    para(doc,'SCRIBER connects field and healthcare facility reporting with a web command dashboard through a shared backend. This document explains how the clients, API and Supabase services exchange data, and how local queuing supports reporting when a connection is unavailable.')
    heading(doc,'System architecture diagram')
    picture(doc,'architecture','Figure 1  SCRIBER application components and communication paths')
    para(doc,'The web dashboard supports administration and review. The mobile app supports data capture and report tracking. Both send authenticated requests to the Express API, which coordinates operational records and attachments with Supabase. Dashed connections identify local persistence and event delivery.')
    next_page(doc,'System components and responsibilities')
    table(doc,['Component','Responsibility'],[
      ('Web dashboard','HTML, CSS and JavaScript interface for incident setup, account administration, casualty review, facilities, analytics and exports.'),
      ('Mobile app and PWA','Expo and React Native interface for responders and documenters. Provides incident selection, role-aware entry, records, attachments and submission status.'),
      ('Express API','TypeScript routes, controllers and services validate requests, apply authorization rules, manage records, generate exports and record audit activity.'),
      ('Supabase Auth','Maintains authentication identity and sessions. API middleware validates bearer tokens and loads the active application profile.'),
      ('PostgreSQL','Stores incidents, casualty identity, incident reports, operational events, clinical records, histories and administrative data.'),
      ('Supabase Storage','Stores attachment objects. The API associates object metadata with reports and supplies attachment access to authorized clients.'),
      ('Supabase Realtime','Delivers database change events to dashboard subscriptions so relevant views can refresh. Publication membership and row policies control availability and visibility.'),
      ('Device persistence','AsyncStorage holds cached incident choices and queued submissions. Session storage depends on the platform, with SecureStore used where supported.'),
    ],[1.65,5.35])
    heading(doc,'Deployment responsibilities')
    para(doc,'The repository separates the static dashboard, the Expo web export and the Node API service. The clients receive an API base URL; server configuration supplies the Supabase connection settings. Database, authentication, storage and realtime services share the operational backend. Hosting configuration is represented by netlify.toml and render.yaml.')
    next_page(doc,'Operational data flow')
    for title,text in [
      ('Authentication and access','The client obtains a session through the authentication workflow and sends the bearer token with API requests. Middleware checks the token and the active user profile. Role guards and controller scope checks determine which incident and casualty operations the caller can perform.'),
      ('Online casualty submission','A responder selects an active incident and completes the appropriate report fields. The client submits the payload to the API. The creation workflow checks the client record identifier, writes the casualty and incident report data, and records related assessment or transport information where provided.'),
      ('Offline submission and recovery','After incident options have been cached, a user can prepare a report without API connectivity. The local queue retains the payload, owner association, attachment data, attempt count and status. On retry, the client submits queued work and tracks failures or partially completed uploads. Queued work must remain visible until synchronization completes.'),
      ('Verification and dashboard refresh','A submitted report becomes available to authorized reviewers. Review actions change verification status and retain review history. Dashboard subscriptions react to relevant database events, while refresh behavior depends on the configured publications, policies and view logic.'),
      ('Attachments and traceability','Attachment objects and report metadata have separate storage responsibilities. An upload must be associated with the correct incident report. Audit and history records preserve significant actions so reviewers can understand changes over time.'),
    ]:
        doc.add_paragraph(title,'Heading 2');para(doc,text)
    doc.add_paragraph('Implementation references','Heading 2')
    para(doc,'api/src/app.ts; api/src/middleware/auth.ts; api/src/controllers/casualty.controller.ts; website/app.js; mobile/src/offline/casualtyQueue.ts; netlify.toml; render.yaml.')
    save(doc,'SCRIBER-System-Architecture.docx')


CORE=[
('users','Application profiles linked to authentication identities. Stores role, reporting context, active status and administrative ownership information.'),
('incidents','Operational events with incident codes, location, status and ownership. Provides the scope for reports and response activities.'),
('casualties','Person identity and identification details. A person is kept distinct from the reports that describe involvement in an incident.'),
('casualty_incidents','Central report linking a casualty, incident and encoder. Includes a client record identifier, operational status and verification status.'),
('healthcare_facilities','Directory of receiving facilities with location and contact details. Referenced by reports, transport and facility care records.'),
('evacuation_centers','Incident-linked evacuation locations with capacity, contact and location information. Reports can optionally reference a center.'),
('audit_logs','Administrative and operational action records associated with the actor and relevant entity information.'),
('call_down_staff','Staff directory for response mobilization. Entries can link to application users and support call-down records.'),
('form_drafts','User-owned in-progress forms stored as JSON payloads, with form type, title, status and timestamps.'),
('notifications','User-targeted or broadcast messages with notification type, read status and optional related entity information.'),
]
CARE=[
('casualty_triage_assessments','Repeated triage assessments for an incident report, including assessment system, category and recording context.'),
('casualty_transport_records','Movement and transport information for an incident report, with receiving facility and recorded transport details.'),
('casualty_treatments','Treatment and stabilization information linked to an incident report.'),
('facility_encounters','Arrival and emergency department encounter information, disposition and additional hospital care timestamps.'),
('clinical_procedures','Procedure type and timing for an incident report, with an optional facility link.'),
('icu_encounters','ICU admission, discharge and ventilation information for an incident report, with an optional facility link.'),
('casualty_outcomes','Final outcome details for an incident report. Source declarations specify a unique report link.'),
('casualty_case_links','Groups related role-specific reports using case_id. Source uniqueness covers a report link and the case and role pair.'),
('casualty_status_history','Operational status transitions for an incident report, including the actor, reason and time where declared.'),
('casualty_verification_history','Review decisions and verification changes associated with an incident report and reviewer.'),
('attachments','Report-linked file metadata and storage path. The attachment object itself is held in object storage.'),
]
OPS=[
('incident_response_timelines','Incident response timestamps and timeline information. Source declarations specify a unique incident link.'),
('dmmp_staff_call_downs','Mobilization and staffing records for an incident, including links to staff or users where declared.'),
('ems_vehicle_arrivals','Emergency medical service vehicle arrival records for an incident.'),
('medical_coordination_assessments','Incident-level evaluation of medical coordination. Source declarations specify a unique incident link.'),
('continuity_of_care_assessments','Incident-level disruption and continuity assessment. Source declarations specify a unique incident link.'),
('facility_resource_snapshots','Time-specific resource and capacity observations associated with an incident and an optional facility.'),
('sitreps','Incident situation reports with reporting period, summary and preparation or approval information.'),
('responder_safety_reports','Incident-linked responder safety reporting information.'),
('responder_safety_responses','Responder-specific incident safety responses, role or function, arrival and protective equipment timestamps.'),
('incident_ic_backups','Designated backup incident commanders and priority for an incident.'),
('incident_ic_delegations','Temporary acting incident commander assignments, delegation and resumption details. Permanent incident ownership remains separate.'),
]


def schema_doc():
    catalog=json.loads((HERE/'schema-catalog.json').read_text(encoding='utf-8'))
    assert {x['name'] for x in catalog}=={name for name,_ in CORE+CARE+OPS}
    doc=style_doc('SCRIBER Database Schema','Entity relationships and table descriptions')
    para(doc,'SCRIBER uses a relational data model centered on casualty_incidents. This report record connects a casualty to an incident and supports repeated assessments, transport, facility care, attachments and review history. The following diagrams and table descriptions explain the 32 entities declared in the project SQL.')
    heading(doc,'Core entity relationship diagram')
    picture(doc,'core-erd','Figure 1  Core entities and selected foreign key relationships')
    para(doc,'One incident can contain many reports, and one casualty can be referenced by multiple incident reports. Each report has an encoder and can reference a receiving facility or evacuation center. Child records preserve the sequence of assessments and care events. PK identifies a primary key, FK a foreign key, and UQ a unique constraint.')
    next_page(doc,'Extended schema relationships')
    para(doc,'The extensions separate report-level care history from incident-level operational information. This keeps repeated clinical events distinct from staffing, resource and coordination observations.')
    picture(doc,'schema-extensions','Figure 2  Report extensions and incident operational entities')
    para(doc,'casualty_outcomes and casualty_case_links declare at most one linked row per report in the reviewed definitions. Several incident assessment tables declare a unique incident link. User-linked support entities provide drafts, notifications, audit activity and staff information. The table descriptions on the following pages explain each entity individually.')
    next_page(doc,'Core and support table descriptions')
    table(doc,['Table','Description'],CORE,[2.15,4.85])
    next_page(doc,'Care and report history table descriptions')
    table(doc,['Table','Description'],CARE,[2.15,4.85])
    next_page(doc,'Incident operations table descriptions')
    table(doc,['Table','Description'],OPS,[2.15,4.85])
    next_page(doc,'Identifiers and data integrity')
    for title,text in [
      ('Primary keys and foreign keys','UUID primary keys identify application records. Foreign keys connect reports to incidents, casualties and users, and connect child events to their parent report. Optional links require special attention: deleting a referenced record may cascade, be restricted or clear the reference depending on the specific constraint.'),
      ('Client submission identity','The root schema declares casualty_incidents.client_record_id as unique and required. The creation workflow uses this identifier to recognize an already submitted client record. The database constraint and API behavior must work together when a device retries a submission.'),
      ('Victim codes and linked cases','The later victim-code migration removes global uniqueness for casualties.id_number because codes may repeat across incidents. The API checks codes within an incident. casualty_case_links groups related reports through a UUID case_id; the reviewed SQL does not declare a separate cases table.'),
      ('History and timestamps','Use status and verification histories to understand transitions instead of relying only on the latest value in the report. Assessment, transport and encounter timestamps support reconstruction of the care journey. Created and updated timestamps serve record maintenance purposes and should not substitute for event time.'),
      ('Roles and authorization','Application roles have expanded beyond the original root enum. Role migrations, API guards and deployed database definitions need to agree. Table policies and backend scope checks jointly determine record visibility; administrative server access still requires explicit authorization logic.'),
      ('Schema maintenance','The project contains overlapping CREATE TABLE and ALTER TABLE declarations. Reconcile the live column definitions, indexes, constraints, enum values and row policies before applying a new baseline or migration. The companion schema-review.sql contains read-only queries for this comparison.'),
    ]:
        doc.add_paragraph(title,'Heading 2');para(doc,text)
    doc.add_paragraph('Implementation references','Heading 2')
    para(doc,'UP-Manila-DB.sql; api/sql/*.sql; api/src/controllers/casualty.controller.ts. Exact table declarations and source filenames are retained in schema-catalog.json.')
    save(doc,'SCRIBER-Database-Schema.docx')


def phone_pair(doc):
    # An inline two-column image keeps the two phone layouts aligned in Word.
    from PIL import Image
    a=Image.open(TMP/'mobile-identity.png');b=Image.open(TMP/'mobile-review.png')
    gap=80;combined=Image.new('RGB',(a.width+b.width+gap,max(a.height,b.height)),'white')
    combined.paste(a,(0,0));combined.paste(b,(a.width+gap,0));combined.save(TMP/'mobile-pair.png')
    picture(doc,'mobile-pair','Figure 2  Mobile identity entry on the left and report review on the right',6.3)


def wireframe_doc():
    doc=style_doc('SCRIBER Wireframes','Web and mobile screen layouts and interaction behavior')
    para(doc,'The SCRIBER interface separates command review from field reporting. The web dashboard organizes incidents and incoming reports, while the mobile app guides a responder through entry, review and submission. These wireframes show the layout and explain how users move between screens.')
    heading(doc,'Web dashboard wireframe')
    picture(doc,'web-wireframe','Figure 1  Desktop dashboard with navigation, summary metrics and casualty reports')
    para(doc,'The left navigation provides access to overview, incidents, casualty records, verification and facilities. Summary cards show incident and report counts. The report table presents the record identifier, incident, assessment label, verification state and a detail action. The verification count keeps pending review work visible.')
    para(doc,'A reviewer opens a submitted report to inspect its details and history. Approval changes its status to verified. Returning a report requires a correction reason so the reporting team can understand the requested change.')
    next_page(doc,'Mobile identity entry and submission review')
    para(doc,'Identity entry establishes the active incident and the casualty name or temporary unidentified label. Age is optional. The assessment step follows identity entry; the review screen then brings the entered information together before submission.')
    phone_pair(doc)
    para(doc,'The identity screen groups a small number of fields above the continue action. The review screen repeats the incident, identity, assessment and location, and changes the primary action to match connectivity. Back navigation preserves draft values so the user can correct an earlier step.')
    next_page(doc,'Mobile records and offline queue')
    para(doc,'The Records screen distinguishes work still held on the device from reports already submitted for command review. The connection banner and textual status labels explain the next available action.')
    picture(doc,'mobile-queue','Figure 3  Pending report and submitted reports on the mobile Records screen',2.9)
    para(doc,'A report saved while offline remains pending in the local queue. After reconnecting, Sync all submits queued work and moves completed submissions into the shared register. A report must retain its client identity across retries to avoid duplicate creation. Failed or partially completed work needs a visible recovery action in the operational app.')
    next_page(doc,'Screen descriptions and interaction rules')
    table(doc,['Screen','Purpose and main actions'],[
      ('Web overview','Scan the current response context, synchronized report totals, assessment distribution and latest reports.'),
      ('Web incidents','Select an incident and open its associated reports.'),
      ('Web casualty records','Search by name or identifier and filter by incident or verification status. Open record details and history.'),
      ('Web verification','Inspect submitted reports, approve a record or return it with a required correction reason.'),
      ('Web facilities','Inspect the receiving facility directory and displayed capacity context.'),
      ('Mobile home','See the current incident and queue summary, then begin a casualty report or open Records.'),
      ('Mobile identity','Choose the incident and enter a name or temporary label, optional age and sex information.'),
      ('Mobile assessment','Enter the assessment label, location and report notes before reviewing the report.'),
      ('Mobile review','Confirm the entered details, return to edit, or submit to the server or local queue according to connectivity.'),
      ('Mobile records','Track queued and submitted work, synchronize after reconnecting, and inspect review status and feedback.'),
    ],[1.7,5.3])
    doc.add_paragraph('Validation and accessibility','Heading 2')
    para(doc,'The demonstration requires an active incident and a name or unidentified label. An entered age must fall between 0 and 130. Status is expressed with text as well as color, and assessment colors are user-entered labels. Clear focus indicators, readable field labels and large mobile actions support navigation and touch input.')
    doc.add_paragraph('Implementation references','Heading 2')
    para(doc,'website/index.html; website/app.js; website/styles.css; mobile/src/app; mobile/src/offline/casualtyQueue.ts. The retained review prototype defines the illustrated interactions and uses demonstration records.')
    save(doc,'SCRIBER-Wireframes.docx')


if __name__=='__main__':
    figures();architecture_doc();schema_doc();wireframe_doc()
