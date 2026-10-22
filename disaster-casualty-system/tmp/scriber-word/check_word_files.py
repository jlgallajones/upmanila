from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree

root=Path(__file__).resolve().parents[2]
out=root/'output/docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
specs=[('SCRIBER-System-Architecture.docx',1,'System components and responsibilities'),
       ('SCRIBER-Database-Schema.docx',2,'Incident operations table descriptions'),
       ('SCRIBER-Wireframes.docx',3,'Screen descriptions and interaction rules')]
for name,image_count,required in specs:
    path=out/name;doc=Document(path)
    text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert 'SCRIBER' in doc.core_properties.title and required in text
    assert 'medsort' not in text.lower()
    assert len(doc.inline_shapes)==image_count
    assert sum(p.style.name=='Title' for p in doc.paragraphs)==1
    with ZipFile(path) as z:
        assert z.testzip() is None
        for file in z.namelist():
            if file.endswith('.xml'):etree.fromstring(z.read(file))
        xml=etree.fromstring(z.read('word/document.xml'))
        assert len(xml.xpath('//w:tblHeader',namespaces=ns))==len(doc.tables)
    print(f'{name}: valid OOXML, {len(doc.tables)} tables, {image_count} embedded figures, {len(text.split())} words')
